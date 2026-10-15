"""
Excel export: your data as a spreadsheet you can open on your iPhone or PC.

The file has three sheets:
    Habits      - one row per day (calories, water, workouts, spending...)
                  with green cells for days you hit your goals, plus charts
    Spending    - every expense, totals by type, and a pie chart
    All entries - everything you've logged, with filter buttons on each column

openpyxl is the library that writes .xlsx files (the Excel format).
"""

from collections import defaultdict
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import tracker

HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="5B4FCF")  # purple, like the bot
GREEN = PatternFill("solid", start_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE")


def _goal(name: str) -> float | None:
    for goal in tracker.get_goals():
        if goal["name"] == name:
            return goal["target"]
    return None


def _minutes(entry: dict) -> float:
    """Study/workout time can be logged in minutes or hours; count it in minutes.
    Anything not measured in time (e.g. "3 sets") counts as 0 minutes."""
    amount, unit = entry["amount"] or 0, (entry["unit"] or "").lower()
    if unit in ("h", "hr", "hrs", "hour", "hours"):
        return amount * 60
    return amount if unit in ("m", "min", "mins", "minute", "minutes") else 0


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is not None:
            cell.font, cell.fill = HEADER, HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _widths(sheet, widths: list[int]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _habits_sheet(book: Workbook, entries: list[dict], start: str, end: str) -> None:
    sheet = book.active
    sheet.title = "Habits"
    #           A       B      C                 D             E           F
    columns = ["Date", "Day", "Calories (kcal)", "Water (ml)", "Workouts", "Workout (min)",
               # G               H                 I              J              K
               "Distance (km)", "Burned (kcal)", "Supplements", "Study (min)", "Sleep (h)",
               # L             M             N        O
               "Spent (RM)", "Tasks done", "Steps", "Mood"]
    sheet.append(columns)

    # Add up each day's entries: days["2026-10-13"]["water"] = 1500, etc.
    days = defaultdict(lambda: defaultdict(float))
    moods = {}
    for e in entries:
        day, category = e["logged_at"][:10], e["category"]
        totals = days[day]
        if category in ("food", "drink"):
            totals["kcal"] += e["calories"] or 0
        elif category == "water":
            totals["water"] += e["amount"] or 0
        elif category == "workout":
            totals["workouts"] += 1
            totals["workout_min"] += e.get("duration_min") or _minutes(e)
            totals["km"] += e.get("distance_km") or 0
            totals["burned"] += e.get("burned_kcal") or 0
        elif category == "supplement":
            totals["supplements"] += 1
        elif category == "study":
            totals["study"] += _minutes(e)
        elif category == "sleep":
            totals["sleep"] += e["amount"] or 0
        elif category == "expense":
            totals["spent"] += e["amount"] or 0
        elif category == "task":
            totals["tasks"] += 1
        elif category == "steps":
            totals["steps"] += e["amount"] or 0
        elif category == "mood":
            moods[day] = e["item"]

    # One row per calendar day, including empty days (so gaps are visible).
    day = date.fromisoformat(start)
    last = min(date.fromisoformat(end), tracker.now().date())
    while day <= last:
        key, t = day.isoformat(), days[day.isoformat()]
        sheet.append([day, day.strftime("%a"), round(t["kcal"]) or None, round(t["water"]) or None,
                      int(t["workouts"]) or None, round(t["workout_min"]) or None,
                      round(t["km"], 2) or None, round(t["burned"]) or None,
                      int(t["supplements"]) or None, round(t["study"]) or None,
                      t["sleep"] or None, round(t["spent"], 2) or None, int(t["tasks"]) or None,
                      int(t["steps"]) or None, moods.get(key)])
        day += timedelta(days=1)
    first_row, last_row = 2, sheet.max_row

    # Totals and averages under the table. These are real Excel formulas,
    # so they update if you edit numbers in the sheet yourself.
    sheet.append([])
    sheet.append(["Average / day", None] + [f"=IFERROR(ROUND(AVERAGE({c}{first_row}:{c}{last_row}),1),\"\")"
                                            for c in "CDEFGHIJKLMN"])
    sheet.append(["Total", None] + [f"=SUM({c}{first_row}:{c}{last_row})" for c in "CDEFGHIJKLMN"])
    for row in (sheet.max_row - 1, sheet.max_row):
        for cell in sheet[row]:
            cell.font = Font(bold=True)

    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, max_col=1):
        row[0].number_format = "dd mmm yyyy"
    _style_header(sheet)
    _widths(sheet, [13, 6, 11, 10, 10, 10, 10, 10, 12, 10, 9, 10, 10, 9, 14])
    sheet.freeze_panes = "C2"  # keep dates and headers visible while scrolling

    # Colour the cells: green = goal met / habit done, red = over the calorie goal.
    data = lambda c: f"{c}{first_row}:{c}{last_row}"
    water_goal, kcal_goal = _goal("water"), _goal("calories")
    if water_goal:
        sheet.conditional_formatting.add(data("D"), CellIsRule(operator="greaterThanOrEqual",
                                                               formula=[str(water_goal)], fill=GREEN))
    if kcal_goal:
        sheet.conditional_formatting.add(data("C"), CellIsRule(operator="greaterThan",
                                                               formula=[str(kcal_goal)], fill=RED))
    for column in "EIJM":  # workouts, supplements, study, tasks: any amount counts as done
        sheet.conditional_formatting.add(data(column), FormulaRule(formula=[f"{column}{first_row}>0"],
                                                                   fill=GREEN))

    # Charts, placed to the right of the table.
    dates = Reference(sheet, min_col=1, min_row=first_row, max_row=last_row)
    for anchor, column, title, chart in (("Q2", 3, "Calories per day", LineChart()),
                                         ("Q18", 4, "Water per day (ml)", BarChart()),
                                         ("Q34", 12, "Spending per day (RM)", BarChart())):
        chart.title, chart.height, chart.width, chart.legend = title, 7, 16, None
        chart.add_data(Reference(sheet, min_col=column, min_row=1, max_row=last_row), titles_from_data=True)
        chart.set_categories(dates)
        chart.x_axis.number_format = "dd mmm"
        sheet.add_chart(chart, anchor)


def _spending_sheet(book: Workbook, entries: list[dict]) -> None:
    sheet = book.create_sheet("Spending")
    expenses = [e for e in entries if e["category"] == "expense"]
    sheet.append(["Date", "What", "Type", "RM", "Notes"])
    for e in expenses:
        sheet.append([e["logged_at"], e["item"], e["tag"] or "other", e["amount"], e["notes"]])
    _style_header(sheet)
    _widths(sheet, [17, 24, 12, 10, 30, 3, 14, 12])
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "#,##0.00"
    if not expenses:
        sheet["A2"] = "No spending logged in this period."
        return

    # Totals per type, biggest first, plus a pie chart of them.
    by_type = defaultdict(float)
    for e in expenses:
        by_type[e["tag"] or "other"] += e["amount"] or 0
    sheet["G1"], sheet["H1"] = "Type", "Total (RM)"
    _style_header(sheet)
    for i, (kind, total) in enumerate(sorted(by_type.items(), key=lambda kv: -kv[1]), start=2):
        sheet.cell(i, 7, kind)
        sheet.cell(i, 8, round(total, 2)).number_format = "#,##0.00"
    total_row = len(by_type) + 2
    sheet.cell(total_row, 7, "TOTAL").font = Font(bold=True)
    sheet.cell(total_row, 8, f"=SUM(H2:H{total_row - 1})").font = Font(bold=True)
    sheet.cell(total_row, 8).number_format = "#,##0.00"

    pie = PieChart()
    pie.title, pie.height, pie.width = "Spending by type", 8, 12
    pie.add_data(Reference(sheet, min_col=8, min_row=1, max_row=total_row - 1), titles_from_data=True)
    pie.set_categories(Reference(sheet, min_col=7, min_row=2, max_row=total_row - 1))
    sheet.add_chart(pie, f"G{total_row + 2}")


def _all_entries_sheet(book: Workbook, entries: list[dict]) -> None:
    sheet = book.create_sheet("All entries")
    columns = ["logged_at", "category", "tag", "item", "amount", "unit", "calories",
               "distance_km", "duration_min", "burned_kcal", "notes"]
    sheet.append(["When", "Category", "Tag", "Item", "Amount", "Unit", "kcal eaten",
                  "km", "min", "kcal burned", "Notes"])
    for e in entries:
        sheet.append([e.get(c) or None for c in columns])
    _widths(sheet, [17, 12, 12, 26, 9, 8, 10, 7, 7, 11, 40])
    if entries:
        # An Excel "table" gives every column a filter/sort button.
        table = Table(displayName="Entries", ref=f"A1:K{len(entries) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "A2"


def build_excel(start: str, end: str) -> BytesIO | None:
    """Make the .xlsx file in memory (no file left on disk). None if there's no data."""
    entries = tracker.entries_between(start, end)
    if not entries:
        return None
    book = Workbook()
    _habits_sheet(book, entries, start, end)
    _spending_sheet(book, entries)
    _all_entries_sheet(book, entries)
    file = BytesIO()
    book.save(file)
    file.seek(0)
    return file
